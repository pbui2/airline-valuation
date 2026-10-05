import json
import shutil
import sys
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from airline_valuation.workbook import Workbook, WorkbookContractError

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "fcffsimpleginzu.xlsx"
CONTRACT = ROOT / "src/airline_valuation/contract.json"
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def rewrite(source, target, changes):
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename in changes:
                data = changes[info.filename](data)
            dst.writestr(info, data)


def xml_change(data, callback):
    root = ET.fromstring(data)
    callback(root)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def test_metadata_fingerprint_is_stable_after_openpyxl_import():
    import openpyxl  # noqa: F401 - registers process-global ElementTree namespaces

    with Workbook(TEMPLATE, CONTRACT) as book:
        assert book.verify_structure()["valid"]


def test_real_contract_and_explicit_input_allowlist():
    with Workbook(TEMPLATE, CONTRACT) as book:
        assert book.verify_identity()["matches_template"]
        assert book.verify_structure()["valid"]
        assert len(book._formula_fingerprints) == 17
        for key in book.contract["input_allowlist"]:
            sheet, address = key.split("!")
            assert book._cells[sheet][address]["formula"] is None
        for cell in ("B27", "B28", "B32", "B35", "B46", "B8", "B9"):
            with pytest.raises(PermissionError):
                book.read_input(cell)
        status = book.branch_status()
        assert not status["ready"]
        assert status["unsupported_branches"] == ["Cost of capital worksheet!B11"]


@pytest.mark.parametrize("mutation,match", [
    ("name", "defined names changed"),
    ("hidden", "sheet state changed"),
    ("validation", "worksheet metadata changed"),
    ("protection", "worksheet metadata changed"),
])
def test_names_sheet_state_and_validation_protection_are_guarded(tmp_path, mutation, match):
    target = tmp_path / f"{mutation}.xlsx"
    def workbook_xml(data):
        def edit(root):
            if mutation == "name":
                names = root.find("m:definedNames", NS)
                if names is None:
                    names = ET.SubElement(root, f"{{{NS['m']}}}definedNames")
                ET.SubElement(names, f"{{{NS['m']}}}definedName", {"name":"AddedName"}).text = "'Input sheet'!$B$11"
            elif mutation == "hidden":
                root.find("m:sheets/m:sheet", NS).set("state", "hidden")
        return xml_change(data, edit)
    def worksheet_xml(data):
        def edit(root):
            if mutation == "validation":
                dv = root.find("m:dataValidations/m:dataValidation", NS)
                assert dv is not None
                dv.set("sqref", "B23 B24")
            elif mutation == "protection":
                ET.SubElement(root, f"{{{NS['m']}}}sheetProtection", {"sheet":"1", "objects":"1"})
        return xml_change(data, edit)
    changes = {}
    if mutation in ("name", "hidden"):
        changes["xl/workbook.xml"] = workbook_xml
    else:
        changes["xl/worksheets/sheet1.xml"] = worksheet_xml
    rewrite(TEMPLATE, target, changes)
    with Workbook(target, CONTRACT) as book:
        with pytest.raises(WorkbookContractError, match=match):
            book.verify_structure()


def test_formula_and_static_mutations_rejected(tmp_path):
    target = tmp_path / "formula.xlsx"
    rewrite(TEMPLATE, target, {"xl/worksheets/sheet2.xml": lambda b: b.replace(b"B31/B32", b"B31/B33", 1)})
    with Workbook(target, CONTRACT) as book, pytest.raises(WorkbookContractError, match="formula"):
        book.verify_structure()
    target = tmp_path / "label.xlsx"
    rewrite(TEMPLATE, target, {"xl/sharedStrings.xml": lambda b: b.replace(b"Current stock price", b"Changed stock price", 1)})
    with Workbook(target, CONTRACT) as book, pytest.raises(WorkbookContractError, match="static cell inventory"):
        book.verify_structure()


def test_feature_content_mutation_and_removal_rejected(tmp_path):
    target = tmp_path / "drawing.xlsx"
    rewrite(TEMPLATE, target, {"xl/drawings/drawing1.xml": lambda b: b + b" "})
    with Workbook(target, CONTRACT) as book, pytest.raises(WorkbookContractError, match="feature content"):
        book.verify_structure()
    target = tmp_path / "table.xlsx"
    rewrite(TEMPLATE, target, {"xl/tables/table1.xml": lambda b: b""})
    with Workbook(target, CONTRACT) as book, pytest.raises(WorkbookContractError, match="feature"):
        book.verify_structure()


def test_excel_calcid_and_calcchain_variation_is_permitted_but_iteration_is_not(tmp_path):
    target = tmp_path / "excel-save.xlsx"
    def workbook(data): return data.replace(b'calcId="191029"', b'calcId="999999"')
    def rels(data):
        def edit(root):
            for r in list(root):
                if r.attrib.get("Type", "").endswith("/calcChain"): root.remove(r)
        return xml_change(data, edit)
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename == "xl/calcChain.xml": continue
            data = src.read(info.filename)
            if info.filename == "xl/workbook.xml": data = workbook(data)
            if info.filename == "xl/_rels/workbook.xml.rels": data = rels(data)
            dst.writestr(info, data)
    with Workbook(target, CONTRACT) as book: assert book.verify_structure()["valid"]
    target = tmp_path / "iterate.xlsx"
    rewrite(TEMPLATE, target, {"xl/workbook.xml": lambda b: b.replace(b'calcId="191029"', b'calcId="191029" iterate="1"')})
    with Workbook(target, CONTRACT) as book, pytest.raises(WorkbookContractError, match="iterate"):
        book.verify_structure()


def test_unsupported_and_unknown_switches_fail_closed(tmp_path):
    for value, expected in (("Industry Average", "unsupported_branches"), ("Novel", "unresolved_switches")):
        target = tmp_path / f"{value}.xlsx"
        def mutate(data):
            def edit(root):
                cell = root.find(".//m:c[@r='B11']", NS)
                cell.set("t", "str")
                cell.find("m:v", NS).text = value
            return xml_change(data, edit)
        rewrite(TEMPLATE, target, {"xl/worksheets/sheet10.xml": mutate})
        with Workbook(target, CONTRACT) as book:
            assert not book.branch_status()["ready"]
            assert book.branch_status()[expected] == ["Cost of capital worksheet!B11"]


def test_cached_output_rejects_boolean_nonfinite_and_error_values(tmp_path):
    target = tmp_path / "boolean.xlsx"
    def update(data, value, typ):
        def edit(root):
            c = root.find(".//m:c[@r='B33']", NS); c.set("t", typ); c.find("m:v", NS).text = value
        return xml_change(data, edit)
    rewrite(TEMPLATE, target, {"xl/worksheets/sheet2.xml": lambda b: update(b, "1", "b")})
    with Workbook(target, CONTRACT) as book:
        assert "Valuation output!B33" in book.branch_status()["type_failures"] if book.branch_status()["branch"] == "direct_input" else not book.branch_status()["ready"]


def test_conditional_inputs_activate_only_when_override_yes(tmp_path):
    with Workbook(TEMPLATE, CONTRACT) as book:
        # Existing template has B45=No; B46 must not be required.
        assert "Input sheet!B46" not in book.branch_status()["missing_required_inputs"]


def test_saved_mode_allows_only_inventoried_formula_followers(tmp_path):
    target = tmp_path / "saved-follower.xlsx"
    def follower_change(data):
        def edit(root):
            cell = root.find(".//m:c[@r='J6']", NS)
            if cell is None:
                row = root.find(".//m:sheetData/m:row[@r='6']", NS)
                cell = ET.SubElement(row, f"{{{NS['m']}}}c", {"r": "J6", "t": "s"})
                ET.SubElement(cell, f"{{{NS['m']}}}v").text = "0"
            else:
                cell.set("t", "s")
                value = cell.find("m:v", NS)
                if value is None: value = ET.SubElement(cell, f"{{{NS['m']}}}v")
                value.text = "0"
        return xml_change(data, edit)
    rewrite(TEMPLATE, target, {"xl/worksheets/sheet10.xml": follower_change})
    with Workbook(target, CONTRACT) as book:
        assert {"J6", "J12"} <= book._formula_followers["Cost of capital worksheet"]
        assert book._shared_formula_groups["Cost of capital worksheet"][0]["range"] == "J5:J12"
        assert book.verify_structure(mode="saved")["valid"]
        with pytest.raises(WorkbookContractError, match="static cell inventory"):
            book.verify_structure(mode="strict")
    unrelated = tmp_path / "saved-static.xlsx"
    rewrite(TEMPLATE, unrelated, {"xl/worksheets/sheet1.xml": lambda data: data.replace(b'<c r="B10" s="344" t="s"><v>883</v></c>', b'<c r="B10" s="344" t="s"><v>0</v></c>', 1)})
    with Workbook(unrelated, CONTRACT) as book:
        with pytest.raises(WorkbookContractError, match="static cell inventory"):
            book.verify_structure(mode="saved")


def test_reader_rejects_duplicate_zip_names_and_traversal(tmp_path):
    duplicate = tmp_path / "duplicate.xlsx"
    with zipfile.ZipFile(TEMPLATE) as src:
        members = [(info, src.read(info.filename)) for info in src.infolist()]
    with zipfile.ZipFile(duplicate, "w") as dst:
        for info, data in members:
            dst.writestr(info, data)
        dst.writestr("xl/workbook.xml", dict((info.filename, data) for info, data in members)["xl/workbook.xml"])
    with pytest.raises(WorkbookContractError, match="duplicate ZIP"):
        Workbook(duplicate, CONTRACT)
    traversal = tmp_path / "traversal.xlsx"
    with zipfile.ZipFile(TEMPLATE) as src, zipfile.ZipFile(traversal, "w") as dst:
        for info in src.infolist(): dst.writestr(info, src.read(info.filename))
        dst.writestr("xl/../evil.xml", b"bad")
    with pytest.raises(WorkbookContractError, match="unsafe ZIP"):
        Workbook(traversal, CONTRACT)


@pytest.mark.parametrize("part,mutate,match", [
    ("_rels/.rels", lambda b: b.replace(b"xl/workbook.xml", b"xl/other.xml", 1), "officeDocument"),
    ("[Content_Types].xml", lambda b: b.replace(b"spreadsheetml.sheet.main+xml", b"text/plain", 1), "workbook content type"),
    ("[Content_Types].xml", lambda b: b.replace(b"worksheet+xml", b"text/plain", 1), "disallowed content type|worksheet content type|content type declarations changed"),
])
def test_reader_rejects_root_relationship_and_content_type_tampering(tmp_path, part, mutate, match):
    target = tmp_path / "package-tamper.xlsx"
    rewrite(TEMPLATE, target, {part: mutate})
    with pytest.raises(WorkbookContractError, match=match):
        Workbook(target, CONTRACT)


def test_saved_mode_rejects_follower_structure_mutation(tmp_path):
    target = tmp_path / "follower-marker.xlsx"
    def mutate(data):
        root = ET.fromstring(data)
        cell = root.find(".//m:c[@r='J6']", NS)
        assert cell is not None
        formula = cell.find("m:f", NS)
        assert formula is not None
        cell.remove(formula)
        value = cell.find("m:v", NS)
        assert value is not None
        value.text = "2"
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)
    rewrite(TEMPLATE, target, {"xl/worksheets/sheet10.xml": mutate})
    with Workbook(target, CONTRACT) as book, pytest.raises(WorkbookContractError, match="follower structure"):
        book.verify_structure(mode="saved")


def test_reader_rejects_bad_relationship_target_and_duplicate_cells(tmp_path):
    target = tmp_path / "bad-target.xlsx"
    rewrite(TEMPLATE, target, {"xl/_rels/workbook.xml.rels": lambda b: b.replace(b"worksheets/sheet1.xml", b"../../outside.xml", 1)})
    with pytest.raises(WorkbookContractError, match="escapes package"):
        Workbook(target, CONTRACT)
    target = tmp_path / "duplicate-cell.xlsx"
    def duplicate_cell(data):
        root = ET.fromstring(data)
        row = root.find(".//m:sheetData/m:row", NS)
        row.append(ET.fromstring(ET.tostring(row.find("m:c", NS))))
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)
    rewrite(TEMPLATE, target, {"xl/worksheets/sheet1.xml": duplicate_cell})
    with pytest.raises(WorkbookContractError, match="duplicate cell coordinate"):
        Workbook(target, CONTRACT)
