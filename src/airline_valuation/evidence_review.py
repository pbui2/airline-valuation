"""Create a static evidence and review workbook; never calculate valuation outputs."""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PDF_URL = "https://investors.southwest.com/sec-filings/all-sec-filings/content/0000092380-26-000004/0000092380-26-000004.pdf"
PDF_SHA256 = "dd8569563afc978e1b61a41396ba07246ac7a85a4545bfda40c58380d4018306"
TREASURY_URL = "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/2026/all?type=daily_treasury_yield_curve"

REPORTED = [
    ["Shares outstanding", "491,317,990", "shares", "Close of business Feb 3, 2026", "PDF p. 1", "Printed page not numbered", "Exact reported cover observation; cutoff eligibility remains a review item."],
    ["Stockholders’ equity", "7,981", "USD millions", "Dec 31, 2025", "PDF p. 87", "Printed p. 84", "Reported balance-sheet line item."],
    ["Current maturities of long-term debt", "324", "USD millions", "Dec 31, 2025", "PDF p. 87", "Printed p. 84", "Reported line item; debt scope requires analyst decision."],
    ["Long-term debt less current maturities", "4,577", "USD millions", "Dec 31, 2025", "PDF p. 87", "Printed p. 84", "Reported line item; debt scope requires analyst decision."],
    ["Current operating lease liabilities", "312", "USD millions", "Dec 31, 2025", "PDF p. 87", "Printed p. 84", "Reported separately; lease treatment unresolved."],
    ["Noncurrent operating lease liabilities", "768", "USD millions", "Dec 31, 2025", "PDF p. 87", "Printed p. 84", "Reported separately; lease treatment unresolved."],
    ["Operating revenues", "28,063", "USD millions", "FY2025", "PDF p. 88", "Printed p. 85", "Reported income-statement line item."],
    ["Operating income", "428", "USD millions", "FY2025", "PDF p. 88", "Printed p. 85", "Reported income-statement line item."],
    ["Interest expense", "167", "USD millions", "FY2025", "PDF p. 88", "Printed p. 85", "Reported income-statement line item; included for reference."],
    ["Income before income taxes", "563", "USD millions", "FY2025", "PDF p. 88", "Printed p. 85", "Reported income-statement line item."],
    ["Provision for income taxes", "122", "USD millions", "FY2025", "PDF p. 88", "Printed p. 85", "Reported income-statement line item."],
]

MAPPING = [
    ["Input sheet!B11", "Historical revenues", "Operating revenues 28,063", "USD millions", "28,063", "Proposed direct mapping; not approved", "FY2025 reported operating revenues; source above, PDF p.88 / printed p.85."],
    ["Input sheet!B12", "Operating income or EBIT", "Operating income 428", "USD millions", "428", "Proposed direct mapping; not approved", "FY2025 reported operating income; source above, PDF p.88 / printed p.85."],
    ["Input sheet!B14", "Book value of equity", "Stockholders’ equity 7,981", "USD millions", "7,981", "Proposed direct mapping; not approved", "Balance-sheet line item; PDF p.87 / printed p.84."],
    ["Input sheet!B15", "Book value of debt", "324 current + 4,577 long-term debt", "USD millions", "4,901 (proposal only)", "NOT APPROVED — debt/lease definition unresolved", "324 + 4,577 = 4,901 is arithmetic on reported components, not a reported subtotal or approved value. Decide whether/how to include operating leases (312 current; 768 noncurrent) and any other obligations."],
    ["Input sheet!B21", "Shares outstanding", "491317990 exact shares", "shares", "491.31799 million shares (proposed conversion)", "NOT APPROVED — conversion/mapping pending", "Keep source observation as exact shares. Division by 1,000,000 is a proposed unit conversion only; reconcile source/cutoff before any mapping."],
    ["Input sheet!B22", "Current stock price", "54.26 USD/share (unverified Feb 6 close)", "USD per share", "—", "PENDING — price/source/rights review", "Source/rights and independent verification unresolved; do not treat as approved."],
    ["Input sheet!B23", "Effective tax rate", "Provision 122; pretax income 563", "USD millions source components", "122 / 563 (proposal only)", "NOT APPROVED — effective-tax convention unresolved", "Quotient is only a proposed convention; not an approved B23 value and not a marginal tax assumption."],
    ["Input sheet!B24", "Marginal tax rate", "No approved observation", "fraction required by contract", "—", "PENDING — analyst assumption", "Source basis, value, rationale, author, and approval required."],
    ["Input sheet!B26", "Next-year revenue growth rate", "No forecast supplied", "fraction required by contract", "—", "PENDING — forecast", "Analyst forecast and approval required."],
    ["Input sheet!B29", "Target pre-tax operating margin", "No forecast supplied", "fraction required by contract", "—", "PENDING — forecast", "Analyst forecast and approval required."],
    ["Input sheet!B30", "Year of margin convergence", "No forecast supplied", "years required by contract", "—", "PENDING — forecast", "Analyst forecast and approval required."],
    ["Input sheet!B31", "Sales to capital ratio, years 1–5", "No forecast supplied", "ratio required by contract", "—", "PENDING — forecast", "Analyst forecast and approval required."],
    ["Input sheet!B34", "Risk-free rate", "10 Yr 4.22% (Treasury row dated Feb 6, 2026)", "percent observation; contract expects fraction", "0.0422 (unit conversion proposal only)", "PENDING — Treasury publication/availability timestamp", "Daily row date does not establish publication before cutoff. No timestamp or source-file hash asserted."],
    ["Cost of capital worksheet!B11", "Cost of capital method selector", "No analyst decision", "choice", "—", "PENDING — direct cost-of-capital selector", "Contract-supported branch is ‘I will input’; selection is an analyst decision and is not made here."],
    ["Cost of capital worksheet!B12", "Direct cost of capital input", "No approved observation", "fraction required by contract", "—", "PENDING — direct cost of capital", "Analyst input, rationale, source/basis, and approval required."],
]


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    ws = workbook.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 48


def build_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me-status"
    readme = [
        ["FY2025 evidence and review workbook", "Southwest Airlines Co."],
        ["Status", "DRAFT — NOT CALCULATED — NOT APPROVED"],
        ["Purpose", "Evidence transcription and review prompts only. This file is not a valuation model, approved input manifest, or authorization to populate the supplied model."],
        ["Model workflow", "Do not use this workbook as input for inspect or prepare. Do not enter or copy values into the valuation workbook based only on this review packet."],
        ["Source", "Southwest Airlines Co. FY2025 Form 10-K PDF (issuer-hosted)"],
        ["Source URL", PDF_URL],
        ["Source SHA-256", PDF_SHA256],
        ["Reported values", "See Verified 10-K for transcribed observations, exact PDF pages, printed pages, units, and qualifications."],
        ["Input mapping", "See Input mapping. Source observations and proposed conversions are in separate columns; proposals are not approved cell values."],
        ["Treasury observation", "Feb 6, 2026: 10 Yr 4.22%; 6 Mo 3.59%. Publication/availability timestamp is unknown; cutoff eligibility pending."],
        ["Treasury source URL", TREASURY_URL],
        ["Unresolved review", "Price and rights; Treasury timestamp; debt/lease definition; effective-tax convention; analyst forecasts; marginal tax; direct cost-of-capital selector and input."],
        ["Arithmetic note", "324 + 4,577 = 4,901 is proposed debt-only arithmetic, not reported subtotal or approved value. 122 / 563 is a proposed effective-tax convention, not an approved rate."],
        ["Workbook behavior", "No model valuation formulas, valuation math, cached outputs, or source model workbook are included. All content is static evidence/review text."],
    ]
    for row in readme:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 115
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 42 if row[0].row > 1 else 28
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=14)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws[2][0].font = Font(bold=True)
    ws[2][1].font = Font(bold=True, color="C00000")
    _sheet(wb, "Verified 10-K", ["Observation", "Reported value", "Source unit", "Period / as of", "Exact PDF page", "Printed page", "Qualification"], REPORTED, [36, 22, 20, 30, 18, 20, 72])
    _sheet(wb, "Input mapping", ["Required contract cell", "Contract meaning", "Source observation (not a cell value)", "Source unit / expected", "Proposed conversion or convention (not approved)", "Review status", "Provenance / decision required"], MAPPING, [32, 35, 46, 31, 42, 50, 88])
    return wb


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate a static FY2025 evidence/review workbook (not a valuation model).")
    parser.add_argument("--output", type=Path, default=Path("examples/fy2025-evidence-review.xlsx"))
    parser.add_argument("--overwrite", action="store_true", help="replace an existing output workbook")
    args = parser.parse_args(argv)
    output = args.output
    if output.exists() and not args.overwrite:
        parser.error(f"refusing to overwrite existing file: {output} (pass --overwrite to replace)")
    output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook().save(output)
    print(f"Created evidence/review workbook: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
